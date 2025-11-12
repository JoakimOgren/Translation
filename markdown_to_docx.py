#!/usr/bin/env python3
"""
Convert Swedish chapter markdown files to a Word docx document.
"""

import os
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import markdown
from html.parser import HTMLParser


class MarkdownToDocxConverter(HTMLParser):
    """Convert HTML (from markdown) to docx format."""

    def __init__(self, doc):
        super().__init__()
        self.doc = doc
        self.current_paragraph = None
        self.current_run = None
        self.bold = False
        self.italic = False
        self.in_heading = False
        self.heading_level = 0

    def handle_starttag(self, tag, attrs):
        if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            self.in_heading = True
            self.heading_level = int(tag[1])
            self.current_paragraph = self.doc.add_heading(level=self.heading_level)
            self.current_paragraph.text = ''
        elif tag == 'p':
            self.current_paragraph = self.doc.add_paragraph()
        elif tag == 'strong' or tag == 'b':
            self.bold = True
        elif tag == 'em' or tag == 'i':
            self.italic = True
        elif tag == 'br':
            if self.current_paragraph:
                self.current_paragraph.add_run('\n')

    def handle_endtag(self, tag):
        if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            self.in_heading = False
            self.heading_level = 0
            self.current_paragraph = None
        elif tag == 'p':
            self.current_paragraph = None
        elif tag == 'strong' or tag == 'b':
            self.bold = False
        elif tag == 'em' or tag == 'i':
            self.italic = False

    def handle_data(self, data):
        if data.strip():  # Only process non-empty data
            if self.current_paragraph is None:
                self.current_paragraph = self.doc.add_paragraph()

            run = self.current_paragraph.add_run(data)
            if self.bold:
                run.bold = True
            if self.italic:
                run.italic = True


def convert_markdown_to_docx_simple(markdown_text, doc):
    """
    Convert markdown text to docx format using a simple approach.

    Args:
        markdown_text: Markdown text to convert
        doc: python-docx Document object
    """
    lines = markdown_text.split('\n')
    current_paragraph = None

    for line in lines:
        # Handle headers
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
        # Handle empty lines
        elif line.strip() == '':
            current_paragraph = None
        # Handle regular text
        else:
            if current_paragraph is None:
                current_paragraph = doc.add_paragraph()
            else:
                current_paragraph.add_run('\n')

            # Simple bold and italic handling
            text = line

            # Process bold/italic markers
            # This is a simple approach - for more complex markdown, use a proper parser
            parts = re.split(r'(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)', text)

            for part in parts:
                if part.startswith('***') and part.endswith('***'):
                    # Bold and italic
                    run = current_paragraph.add_run(part[3:-3])
                    run.bold = True
                    run.italic = True
                elif part.startswith('**') and part.endswith('**'):
                    # Bold
                    run = current_paragraph.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    # Italic
                    run = current_paragraph.add_run(part[1:-1])
                    run.italic = True
                else:
                    current_paragraph.add_run(part)


def convert_chapters_to_docx(chapters_dir, output_file):
    """
    Convert Swedish chapter markdown files to a single Word docx file.

    Args:
        chapters_dir: Directory containing chapter markdown files
        output_file: Output docx file path
    """
    print(f"Converting Swedish chapters to Word document...")
    print(f"Input directory: {chapters_dir}")
    print(f"Output file: {output_file}")
    print("="*60)

    # Get all chapter files in sorted order
    chapter_files = sorted(Path(chapters_dir).glob("chapter_*.md"))

    if not chapter_files:
        print(f"Error: No chapter files found in {chapters_dir}")
        return False

    print(f"Found {len(chapter_files)} chapter files to convert")

    # Create a new Document
    doc = Document()

    # Set document properties
    doc.core_properties.title = "Ministry Fulfillment - Swedish Translation"
    doc.core_properties.author = "Translated by Claude"

    # Process each chapter
    for i, chapter_file in enumerate(chapter_files, 1):
        print(f"\n[{i}/{len(chapter_files)}] Processing: {chapter_file.name}")

        # Read chapter content
        with open(chapter_file, 'r', encoding='utf-8') as f:
            markdown_text = f.read()

        # Add page break between chapters (except before first chapter)
        if i > 1:
            doc.add_page_break()

        # Convert markdown to docx
        convert_markdown_to_docx_simple(markdown_text, doc)

        print(f"  ✓ Converted {len(markdown_text)} characters")

    # Save the document
    doc.save(output_file)

    print("\n" + "="*60)
    print(f"✓ Successfully created Word document: {output_file}")
    print(f"  Total chapters: {len(chapter_files)}")

    return True


def main():
    if len(sys.argv) < 3:
        print("Usage: python markdown_to_docx.py <chapters_dir> <output_file>")
        print("\nExample:")
        print("  python markdown_to_docx.py Swedish/chapters Swedish_Ministry_Book.docx")
        sys.exit(1)

    chapters_dir = sys.argv[1]
    output_file = sys.argv[2]

    if not os.path.isdir(chapters_dir):
        print(f"Error: Directory not found: {chapters_dir}")
        sys.exit(1)

    # Ensure output file has .docx extension
    if not output_file.endswith('.docx'):
        output_file += '.docx'

    success = convert_chapters_to_docx(chapters_dir, output_file)

    if not success:
        sys.exit(1)


if __name__ == "__main__":
    main()
